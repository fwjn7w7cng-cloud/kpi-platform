"""Цикл месяца: открыт → согласование → утверждён.

  open      месяц идёт: отчёты принимаются, табель и правки вносятся
  review    согласование: новые автоматические отчёты за месяц не принимаются, руководители
            проверяют табель и правки, руководитель решает по правкам
  approved  утверждён: всё закрыто, снимок расчёта сохранён, ведомость и листки готовы

Правка зарплаты — премия (+) или удержание (−) с обязательной причиной. С модулем «Цепочка утверждения»
правку предлагает руководитель отдела, а принимает руководитель компании; без модуля правку вносит
сразу руководитель компании или админ.
"""
import io
from datetime import datetime

from core.calendar import month_title

NEXT = {"open": "review", "review": "approved"}
TITLES = {"open": "открыт", "review": "на согласовании", "approved": "утверждён"}


class MonthError(Exception):
    pass


def status(db, month):
    return db.month_status(month)


def can_edit(db, month):
    return status(db, month) != "approved"


def advance(db, month, who):
    """Следующий шаг цикла. Утверждение — со снимком расчёта (передаётся отдельно, см. approve)."""
    st = status(db, month)
    if st not in NEXT:
        raise MonthError(f"{month_title(month)} уже {TITLES[st]}")
    if NEXT[st] == "approved":
        raise MonthError("утверждение — через approve(), со снимком расчёта")
    db.set_month_status(month, NEXT[st], who)
    return NEXT[st]


def approve(db, month, who, payroll):
    """Утвердить: правки, по которым не решили, отклоняются; расчёт всех сохраняется снимком."""
    st = status(db, month)
    if st == "approved":
        raise MonthError(f"{month_title(month)} уже утверждён")
    pending = db.q("SELECT id FROM edits WHERE month = ? AND status = 'wait'", month)
    for e in pending:
        db.x("UPDATE edits SET status = 'no', decided_by = ?, decided_ts = ? WHERE id = ?", f"{who} (при утверждении)", _now(), e["id"])
    snap = {}
    for e in payroll.kpi.employees.values():
        slip = payroll.calc(e)
        if not slip.rows and not slip.total:
            continue
        snap[str(e["id"])] = {"fio": e["fio"], "dept": e["dept"], "position": e["position"], "total": slip.total,
                              "days": slip.days, "rows": [[r[0], round(r[1]), r[2]] for r in slip.rows]}
    db.set_month_status(month, "approved", who, snapshot=snap)
    return snap


def reopen(db, month, who, reason):
    if not reason:
        raise MonthError("открыть утверждённый месяц можно только с причиной")
    db.set_month_status(month, "open", who)
    db.audit(who, "month_reopen", f"{month}: {reason}")


def snapshot(db, month):
    r = db.one("SELECT snapshot FROM months WHERE month = ?", month)
    import json
    return json.loads(r["snapshot"]) if r and r["snapshot"] else None


# ---------- правки ----------

def propose(db, month, emp_id, amount, reason, author, direct=False):
    """Предложить правку. direct=True — сразу принята (без модуля цепочки или от руководителя)."""
    if not can_edit(db, month):
        raise MonthError(f"{month_title(month)} утверждён — правки закрыты")
    reason = (reason or "").strip()
    if len(reason) < 3:
        raise MonthError("у правки обязательно должна быть причина")
    if not amount:
        raise MonthError("сумма правки не может быть нулём")
    st = "ok" if direct else "wait"
    eid = db.x("INSERT INTO edits (month, emp_id, amount, reason, author, status, decided_by, ts, decided_ts) "
               "VALUES (?,?,?,?,?,?,?,?,?)", month, emp_id, amount, reason, author, st,
               author if direct else "", _now(), _now() if direct else None)
    db.audit(author, "edit_add", f"{month} #{eid} {emp_id}: {amount:+.0f} — {reason}")
    return eid


def decide(db, edit_id, ok, who):
    e = db.one("SELECT * FROM edits WHERE id = ?", edit_id)
    if not e:
        raise MonthError("правка не найдена")
    if e["status"] != "wait":
        raise MonthError("по этой правке уже решили")
    if not can_edit(db, e["month"]):
        raise MonthError("месяц утверждён")
    db.x("UPDATE edits SET status = ?, decided_by = ?, decided_ts = ? WHERE id = ?", "ok" if ok else "no", who, _now(), edit_id)
    db.audit(who, "edit_ok" if ok else "edit_no", f"#{edit_id}")


def edits(db, month, status=None):
    q = "SELECT * FROM edits WHERE month = ?" + (" AND status = ?" if status else "") + " ORDER BY id"
    return db.q(q, *((month, status) if status else (month,)))


# ---------- ведомость ----------

def registry_xlsx(cfg, month, slips):
    """Ведомость для бухгалтера: отдел → люди, строки расчёта, итог. slips — [Payslip]."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = "Ведомость"
    ws.append([f"{cfg.name} · ведомость за {month_title(month).lower()}"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append(["Предварительный расчёт. Официальная зарплата — в бухгалтерии."])
    ws.append([])
    head = ["Отдел", "Сотрудник", "Должность", "Дней", "Начислено", "Удержано", "К выплате"]
    ws.append(head)
    for c in ws[4]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="E8EFFD")
    total = 0
    for s in sorted(slips, key=lambda s: ((cfg.dept(s.emp["dept"]).name if cfg.dept(s.emp["dept"]) else ""), s.emp["fio"])):
        if not s.rows:
            continue
        plus = sum(r[1] for r in s.rows if r[1] > 0)
        minus = -sum(r[1] for r in s.rows if r[1] < 0)
        pos = cfg.position(s.emp["position"])
        dept = cfg.dept(s.emp["dept"])
        ws.append([dept.name if dept else s.emp["dept"], s.emp["fio"], pos["name"] if pos else s.emp["position"],
                   s.days, round(plus), round(minus), s.total])
        total += s.total
    ws.append([])
    ws.append(["", "Итого", "", "", "", "", total])
    ws.cell(ws.max_row, 2).font = Font(bold=True)
    ws.cell(ws.max_row, 7).font = Font(bold=True)
    for col, w in zip("ABCDEFG", (16, 34, 26, 7, 13, 12, 13)):
        ws.column_dimensions[col].width = w
    for row in ws.iter_rows(min_row=5, min_col=5, max_col=7):
        for c in row:
            c.number_format = "# ##0"
            c.alignment = Alignment(horizontal="right")

    det = wb.create_sheet("Расчёт по строкам")
    det.append(["Сотрудник", "Строка", "Сумма", "Пояснение"])
    for c in det[1]:
        c.font = Font(bold=True)
    for s in slips:
        for r in s.rows:
            det.append([s.emp["fio"], r[0], round(r[1]), r[2]])
    for col, w in zip("ABCD", (34, 30, 12, 60)):
        det.column_dimensions[col].width = w
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _now():
    return datetime.now().isoformat(timespec="seconds")
