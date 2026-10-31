"""Модуль «Склад для закупок»: остатки по поставщикам, что закончилось и что скоро закончится, что лежит
без движения; заказ поставщику — черновик сам из «нужно», правка текстом, Excel, «в пути», «пришёл».

Данные: последний срез остатков (stock) и продажи по товарам (item_sales).
  скорость    продано за окно stock.window дней (по умолчанию 60) ÷ дней окна
  дней запаса остаток ÷ скорость
  🔴 нет       остатка нет, а за окно товар продавался
  🟠 скоро     запаса меньше stock.low_days дней (по умолчанию 14)
  💤 залежалось остаток есть, продаж нет stock.stale_days дней (по умолчанию 90; если история короче — не считаем)
  нужно       скорость × stock.cover_days (30) − остаток − «в пути» (заказы, отправленные поставщику, до 45 дней)
Видят: руководство, админ и отдел закупок.
"""
import io
import json
from collections import defaultdict
from datetime import date, datetime, timedelta

from core.calendar import day_title
from core.jobs.schedule import job
from core.ui.kit import Button, Reply, money, num, ok_row, rows
from core.ui.registry import free_text, menu, screen, text_input
from core.ui.screens import short
from core.ui.views import Card, Cell, Row, View

M = "stock"
CAT = {"out": ("Закончилось", "bad"), "low": ("Скоро закончится", "warn"), "stale": ("Залежалось", None)}
TRANSIT_DAYS = 45


def ok(ctx):
    return ctx.role in ("owner", "admin") or ctx.dept_kind == "purchasing"


def boss(ctx):
    """Видит всех поставщиков и согласует: директор, админ, начальник закупок."""
    return ctx.role in ("owner", "admin") or (ctx.role == "head" and ctx.dept_kind == "purchasing")


def owners(db):
    """Поставщик → id менеджера по закупкам (кто за кого отвечает)."""
    return {k: int(v) for k, v in (db.get("stock_owners") or {}).items()}


def mine(ctx, supplier):
    if boss(ctx):
        return True
    own = owners(ctx.db)
    return not own or (ctx.me is not None and own.get(supplier) == ctx.me["id"])


def limit(ctx):
    return float((ctx.cfg.raw.get("stock") or {}).get("approve_limit") or 300000)


menu("sk", "📦 Склад", {"owner", "admin", "head", "staff"}, module=M, order=30, when=ok)
menu("sow", "✅ Согласовать заказы", {"owner", "admin", "head"}, module=M, order=29,
     when=lambda c: boss(c) and bool(waiting(c)))
menu("sol", "🛒 Заказы поставщикам", {"head", "staff"}, module=M, order=31, when=lambda c: c.dept_kind == "purchasing")


def _cfg(ctx):
    s = ctx.cfg.raw.get("stock") or {}
    return {k: int(s.get(k, d)) for k, d in (("window", 60), ("low_days", 14), ("stale_days", 90), ("cover_days", 30))}


class Stock:
    def __init__(self, ctx):
        self.ctx = ctx
        c = _cfg(ctx)
        self.c = c
        r = ctx.db.one('SELECT MAX("on") AS d FROM stock WHERE "on" <= ?', ctx.today)
        self.on = r["d"] if r and r["d"] else None
        self.items = {x["code"]: x for x in ctx.db.q('SELECT * FROM stock WHERE "on" = ?', self.on)} if self.on else {}
        start = (ctx.today - timedelta(days=c["window"])).isoformat()
        self.sold = defaultdict(float)
        self.amount = defaultdict(float)
        for x in ctx.db.q("SELECT code, SUM(qty) AS q, SUM(amount) AS a FROM item_sales WHERE day > ? AND day <= ? GROUP BY code",
                          start, ctx.today):
            self.sold[x["code"]], self.amount[x["code"]] = x["q"], x["a"]
        first = ctx.db.one("SELECT MIN(day) AS d FROM item_sales")
        self.history = (ctx.today - date.fromisoformat(first["d"])).days if first and first["d"] else 0
        self.last = {x["code"]: x["d"] for x in ctx.db.q("SELECT code, MAX(day) AS d FROM item_sales WHERE qty > 0 GROUP BY code")}
        self.days = min(c["window"], max(self.history, 1))
        self.transit = defaultdict(float)
        since = (ctx.today - timedelta(days=TRANSIT_DAYS)).isoformat()
        for o in ctx.db.q("SELECT items FROM stock_orders WHERE status = 'sent' AND sent_ts >= ?", since):
            for code, q in json.loads(o["items"]):
                self.transit[code] += q

    def speed(self, code):
        return self.sold.get(code, 0) / self.days

    def days_left(self, code):
        s = self.speed(code)
        return self.items[code]["qty"] / s if s > 0 else None

    def price(self, code):
        q = self.sold.get(code, 0)
        return self.amount.get(code, 0) / q if q else 0

    def cat(self, code):
        it = self.items[code]
        if it["qty"] <= 0 and self.sold.get(code, 0) > 0:
            return "out"
        d = self.days_left(code)
        if d is not None and d < self.c["low_days"]:
            return "low"
        if it["qty"] > 0 and self.history >= self.c["stale_days"]:
            last = self.last.get(code)
            if not last or (self.ctx.today - date.fromisoformat(last)).days >= self.c["stale_days"]:
                return "stale"
        return None

    def need(self, code, transit=True):
        n = self.speed(code) * self.c["cover_days"] - max(self.items[code]["qty"], 0) - (self.transit.get(code, 0) if transit else 0)
        return max(round(n), 0)

    def suppliers(self):
        out = defaultdict(list)
        for code, it in self.items.items():
            out[it["supplier"] or "Без поставщика"].append(code)
        return sorted(out.items(), key=lambda t: t[0])

    def my_suppliers(self, ctx):
        """[(номер, поставщик, коды)] — только свои у менеджера, все у начальника и директора."""
        return [(i, sup, codes) for i, (sup, codes) in enumerate(self.suppliers()) if mine(ctx, sup)]


def load(ctx):
    return ctx.app.memo(ctx, "stock", lambda: Stock(ctx), month="-")


def _unit(it):
    return it.get("unit") or "шт"


@screen("sk", module=M)
def stock(ctx):
    if not ok(ctx):
        return Reply(text="Склад видят руководство и отдел закупок.")
    s = load(ctx)
    if not s.on:
        return Reply(text="Остатков пока нет — пришлите отчёт 1С «Остатки товаров».")
    my = s.my_suppliers(ctx)
    codes_all = [c for _, _, cs in my for c in cs]
    cats = defaultdict(list)
    for code in codes_all:
        c = s.cat(code)
        if c:
            cats[c].append(code)
    who = "все поставщики" if boss(ctx) or not owners(ctx.db) else f"мои поставщики: {len(my)}"
    v = View("📦 Склад", f"остатки на {day_title(date.fromisoformat(s.on))} · {who} · товаров {len(codes_all)}")
    stale_sum = sum(s.items[c]["qty"] * (s.price(c) or 0) for c in cats["stale"])
    v.tiles([("Закончилось", str(len(cats["out"])), "bad" if cats["out"] else "good"),
             ("Скоро закончится", str(len(cats["low"])), "warn" if cats["low"] else "good"),
             ("Залежалось", str(len(cats["stale"])) if s.history >= s.c["stale_days"] else "—", None)])
    names = {e["id"]: short(e["fio"]) for e in ctx.db.employees(active=None)}
    own = owners(ctx.db)
    rs = []
    for i, sup, codes in my:
        out_ = sum(1 for c in codes if s.cat(c) == "out")
        low = sum(1 for c in codes if s.cat(c) == "low")
        need = sum(1 for c in codes if s.need(c) > 0)
        rs.append(Row([sup, str(len(codes)), Cell(str(out_), "bad" if out_ else None), Cell(str(low), "warn" if low else None),
                       str(need)], sub=f"менеджер {names[own[sup]]}" if boss(ctx) and own.get(sup) in names else ""))
    v.table(["Поставщик", "Товаров", "Нет", "Скоро", "Заказать"], rs, title="По поставщикам")
    if stale_sum:
        v.note(f"В залежалом товаре — около {money(stale_sum)} по цене продажи.")
    n_wait = len(waiting(ctx)) if boss(ctx) else 0
    kb = [[Button(f"🔴 Нет ({len(cats['out'])})", "skf:out"), Button(f"🟠 Скоро ({len(cats['low'])})", "skf:low")],
          [Button(f"💤 Залежалось ({len(cats['stale'])})", "skf:stale"), Button("🛒 Заказы", "sol")]]
    if n_wait:
        kb.append([Button(f"✅ Согласовать заказы ({n_wait})", "sow")])
    if boss(ctx):
        kb.append([Button("👥 Кто за каких поставщиков", "sku")])
    kb += rows([Button(sup[:24], f"sks:{i}") for i, sup, _ in my], 2)
    v.note("Поиск товара: напишите его код или часть названия.", where="all")
    return Reply(view=v, buttons=kb)


def _item_card(s, code):
    it = s.items[code]
    d = s.days_left(code)
    c = s.cat(code)
    lines = [f"остаток {num(it['qty'])} {_unit(it)} · продаётся {num(round(s.speed(code) * 30))} в месяц",
             f"запаса {'на ' + num(round(d)) + ' дн.' if d is not None else '—'} · нужно заказать {num(s.need(code))}"]
    if s.transit.get(code):
        lines.append(f"в пути {num(s.transit[code])}")
    return Card(f"{it['name']}", f"код {code}", lines, color=CAT[c][1] if c else None)


@screen("skf", module=M)
def by_cat(ctx, cat):
    if not ok(ctx) or cat not in CAT:
        return Reply(text="Нет доступа.")
    s = load(ctx)
    codes = [c for _, _, cs in s.my_suppliers(ctx) for c in cs if s.cat(c) == cat]
    if cat == "stale":
        codes.sort(key=lambda c: -s.items[c]["qty"])
    else:
        codes.sort(key=lambda c: -s.speed(c))
    v = View(f"📦 {CAT[cat][0]}", f"{len(codes)} товаров")
    v.cards([_item_card(s, c) for c in codes[:15]])
    if not codes:
        v.note("Таких товаров нет 👍")
    return Reply(view=v, buttons=[[Button("‹ Склад", "sk")]])


@screen("sks", module=M)
def supplier(ctx, i):
    if not ok(ctx):
        return Reply(text="Нет доступа.")
    s = load(ctx)
    sups = s.suppliers()
    if int(i) >= len(sups) or not mine(ctx, sups[int(i)][0]):
        return stock(ctx)
    sup, codes = sups[int(i)]
    order_ = {"out": 0, "low": 1, None: 2, "stale": 3}
    codes = sorted(codes, key=lambda c: (order_[s.cat(c)], -s.need(c)))
    v = View(f"📦 {sup}", f"товаров {len(codes)} · к заказу {sum(1 for c in codes if s.need(c) > 0)}")
    rs = []
    for c in codes[:40]:
        it, d, ct = s.items[c], s.days_left(c), s.cat(c)
        rs.append(Row([it["name"][:34], num(it["qty"]), Cell(num(round(d)) if d is not None else "—", CAT[ct][1] if ct else None),
                       num(s.need(c))], sub=f"код {c}" + (f" · в пути {num(s.transit[c])}" if s.transit.get(c) else "")))
    v.table(["Товар", "Остаток", "Дней", "Нужно"], rs)
    if len(codes) > 40:
        v.note(f"Показаны 40 из {len(codes)}, сначала срочные.")
    return Reply(view=v, buttons=[[Button("🛒 Заказ этому поставщику", f"son:{i}")], [Button("‹ Склад", "sk")]])


@screen("sku", module=M)
def assign(ctx, i=None, emp=None):
    """Кто из менеджеров отвечает за какого поставщика — назначает начальник закупок или директор."""
    if not boss(ctx):
        return Reply(text="Назначает начальник закупок или директор.")
    s = load(ctx)
    sups = s.suppliers()
    buyers = [e for e in ctx.db.employees() if ctx.cfg.dept(e["dept"]) and ctx.cfg.dept(e["dept"]).kind == "purchasing"
              and e["role"] == "staff"]
    own = owners(ctx.db)
    if i is not None and emp is not None and int(i) < len(sups):
        own = {k: v for k, v in own.items()}
        own[sups[int(i)][0]] = int(emp)
        ctx.db.put("stock_owners", own)
        ctx.db.audit(ctx.author(), "stock_owner", f"{sups[int(i)][0]} → {ctx.db.employee(int(emp))['fio']}")
        i = None
    names = {e["id"]: e["fio"] for e in ctx.db.employees(active=None)}
    if i is not None and int(i) < len(sups):
        sup = sups[int(i)][0]
        return Reply(text=f"Кто отвечает за «{sup}»?",
                     buttons=rows([Button(("• " if own.get(sup) == e["id"] else "") + short(e["fio"]), f"sku:{i}:{e['id']}")
                                   for e in buyers], 2) + [[Button("‹ Назад", "sku")]])
    v = View("👥 Поставщики по менеджерам", f"менеджеров {len(buyers)} · поставщиков {len(sups)}")
    v.table(["Поставщик", "Менеджер"], [Row([sup, short(names[own[sup]]) if own.get(sup) in names else Cell("не назначен", "warn")])
                                       for sup, _ in sups])
    v.note("Менеджер видит в «📦 Склад» и в заказах только своих поставщиков.")
    return Reply(view=v, buttons=rows([Button(sup[:24], f"skm:{n}") for n, (sup, _) in enumerate(sups)], 2))


@screen("skm", module=M)
def assign_one(ctx, i, act=None, value=None):
    """Кто отвечает за поставщика — окно: отметить менеджера, «ОК» / «Отмена»."""
    from core.ui import pick
    if not boss(ctx):
        return Reply(text="Назначает начальник закупок или директор.")
    sups = load(ctx).suppliers()
    if not i.isdigit() or int(i) >= len(sups):
        return assign(ctx)
    sup = sups[int(i)][0]
    buyers = [e for e in ctx.db.employees() if ctx.cfg.dept(e["dept"]) and ctx.cfg.dept(e["dept"]).kind == "purchasing"
              and e["role"] == "staff"]
    own = owners(ctx.db)
    res = pick.one(ctx, f"skm:{i}", act, value, [(str(e["id"]), short(e["fio"])) for e in buyers],
                   str(own.get(sup)) if own.get(sup) else None, f"👥 Кто отвечает за «{sup}»", "отметьте и нажмите «ОК»")
    if res.done and res.value:
        return assign(ctx, i, res.value)
    if res.done or res.cancel:
        return assign(ctx)
    return res.reply


# ---------- заказы ----------
# Путь заказа: черновик → на согласовании → согласован → в пути (отправлен поставщику) → принят складом.
# Согласует начальник закупок, если сумма не больше порога (stock.approve_limit), иначе директор.

ST = {"draft": "черновик", "wait": "⏳ на согласовании", "back": "↩️ возвращён", "ok": "✅ согласован",
      "sent": "🚚 в пути", "came": "📦 принят", "part": "⚠️ принят с недостачей"}


def _order(ctx, oid):
    o = ctx.db.one("SELECT * FROM stock_orders WHERE id = ?", int(oid))
    if o:
        o["items"] = json.loads(o["items"])
        o["extra"] = json.loads(o.get("extra") or "{}")
    return o


def _save(ctx, o, **cols):
    cols["extra"] = json.dumps(o["extra"], ensure_ascii=False)
    sets = ", ".join(f"{k} = ?" for k in cols)
    ctx.db.x(f"UPDATE stock_orders SET {sets} WHERE id = ?", *cols.values(), o["id"])


def _hist(ctx, o, what):
    o["extra"].setdefault("hist", []).append([datetime.now().strftime("%d.%m %H:%M"), ctx.author(), what])


def total(s, o):
    return sum((s.price(c) if c in s.items else 0) * q for c, q in o["items"])


def approver_kind(ctx, o):
    """Кто согласует: начальник закупок (до порога) или директор."""
    return "owner" if total(load(ctx), o) > limit(ctx) else "head"


def can_approve(ctx, o):
    if ctx.role in ("owner", "admin"):
        return True
    return ctx.role == "head" and ctx.dept_kind == "purchasing" and approver_kind(ctx, o) == "head"


def waiting(ctx):
    rs = ctx.db.q("SELECT id FROM stock_orders WHERE status = 'wait'")
    return [r["id"] for r in rs if can_approve(ctx, _order(ctx, r["id"]))]


def _can_edit(ctx, o):
    return o["status"] in ("draft", "back") and (boss(ctx) or (ctx.me and o.get("emp_id") == ctx.me["id"]) or mine(ctx, o["supplier"]))


def _notify(ctx, emps, text, oid):
    return [Reply(text=text, to=e["tg_id"], buttons=[[Button("🛒 Открыть заказ", f"sko:{oid}")]]) for e in emps if e.get("tg_id")]


def _approvers(ctx, o):
    kind = approver_kind(ctx, o)
    out = []
    for e in ctx.db.employees():
        d = ctx.cfg.dept(e["dept"])
        if e["role"] == "owner" or (kind == "head" and e["role"] == "head" and d and d.kind == "purchasing"):
            out.append(e)
    return out


@screen("son", module=M)
def new_order(ctx, i):
    if not ok(ctx):
        return Reply(text="Нет доступа.")
    s = load(ctx)
    sups = s.suppliers()
    if int(i) >= len(sups) or not mine(ctx, sups[int(i)][0]):
        return stock(ctx)
    sup, codes = sups[int(i)]
    draft = ctx.db.one("SELECT id FROM stock_orders WHERE supplier = ? AND status IN ('draft', 'back')", sup)
    if draft:
        return order(ctx, draft["id"])
    items = [[c, s.need(c)] for c in codes if s.need(c) > 0 and s.sold.get(c, 0) > 0]
    oid = ctx.db.x("INSERT INTO stock_orders (supplier, items, author, ts, emp_id) VALUES (?,?,?,?,?)", sup, json.dumps(items),
                   ctx.author(), datetime.now().isoformat(timespec="seconds"), ctx.me["id"] if ctx.me else None)
    return order(ctx, oid)


@screen("sko", module=M)
def order(ctx, oid):
    o = _order(ctx, oid)
    if not o or not (ok(ctx) or ctx.dept_kind == "warehouse") or (ctx.dept_kind == "purchasing" and not mine(ctx, o["supplier"])):
        return Reply(text="Нет доступа.")
    s = load(ctx)
    v = View(f"🛒 Заказ №{o['id']} · {o['supplier']}", f"{ST[o['status']]} · оформил {o['author']}")
    got = dict((c, q) for c, q in o["extra"].get("got", []))
    rs = []
    for code, q in o["items"]:
        it = s.items.get(code, {"name": code, "unit": "шт"})
        p = s.price(code) if code in s.items else 0
        cells = [it["name"][:34], f"{num(q)} {_unit(it)}", money(p * q, "") if p else "—"]
        if o["status"] == "part":
            g = got.get(code, q)
            cells.append(Cell(num(g), "bad" if g < q else "good"))
        rs.append(Row(cells, sub=f"код {code}"))
    head = ["Товар", "Кол-во", "≈ ₽"] + (["Пришло"] if o["status"] == "part" else [])
    v.table(head, rs) if rs else v.note("Пусто. Нажмите «✏️ Изменить» и впишите товары.")
    t = total(s, o)
    if t:
        who = "директор" if t > limit(ctx) else "начальник закупок"
        v.tiles([("Сумма заказа", money(t), None), ("Позиций", str(len(o["items"])), None),
                 ("Согласует", who, "warn" if t > limit(ctx) else None)])
        v.note(f"Сумма — по цене продажи (закупочная ниже). Порог согласования директором — {money(limit(ctx))}.")
    if o["extra"].get("note"):
        v.note(f"Комментарий: {o['extra']['note']}")
    if o["extra"].get("hist"):
        v.lines([f"{ts} · {who_} · {what}" for ts, who_, what in o["extra"]["hist"][-6:]], title="История")
    kb = []
    if _can_edit(ctx, o):
        kb.append([Button("✏️ Изменить", f"soe:{o['id']}"), Button("📨 На согласование", f"sop:{o['id']}")])
    if o["status"] == "wait" and can_approve(ctx, o):
        kb.append([Button("✅ Согласовать", f"soa:{o['id']}"), Button("↩️ Вернуть", f"sob:{o['id']}")])
    if o["status"] == "ok" and (boss(ctx) or mine(ctx, o["supplier"])):
        kb.append([Button("📨 Отправлен поставщику", f"sos:{o['id']}")])
    if o["status"] == "sent" and (ctx.dept_kind == "warehouse" or ctx.role in ("owner", "admin")):
        kb.append([Button("📥 Принять поставку", f"wrc:{o['id']}")])
    kb.append([Button("📥 Excel", f"sox:{o['id']}"), Button("‹ Заказы", "sol")])
    return Reply(view=v, buttons=kb)


@screen("soe", module=M)
def order_edit(ctx, oid):
    o = _order(ctx, oid)
    if not o or not _can_edit(ctx, o):
        return Reply(text="Менять можно черновик или возвращённый заказ.")
    return ctx.ask("stock_ord", "Напишите товары строками «код количество» (0 — убрать):\n\n100012 40\n100031 0",
                   {"id": o["id"]})


@text_input("stock_ord")
def on_edit(ctx, text, data):
    o = _order(ctx, data["id"])
    items = dict((c, q) for c, q in o["items"])
    for line in (text or "").splitlines():
        p = line.replace(",", ".").split()
        if not p:
            continue
        try:
            code, q = p[0], float(p[1]) if len(p) > 1 else None
        except ValueError:
            return ctx.ask("stock_ord", f"Не понял «{line.strip()}». Пишите «код количество».", data)
        if q is None:
            return ctx.ask("stock_ord", f"В строке «{line.strip()}» нет количества.", data)
        if q <= 0:
            items.pop(code, None)
        else:
            items[code] = q
    o["items"] = [[c, q] for c, q in items.items()]
    _hist(ctx, o, "изменён")
    _save(ctx, o, items=json.dumps(o["items"]))
    return order(ctx, o["id"])


@screen("sop", module=M)
def order_propose(ctx, oid):
    o = _order(ctx, oid)
    if not o or not _can_edit(ctx, o):
        return Reply(text="Отправить на согласование можно черновик или возвращённый заказ.")
    if not o["items"]:
        return Reply(text="В заказе нет товаров.")
    t = total(load(ctx), o)
    if can_approve(ctx, o) and ctx.role in ("owner", "admin") or (boss(ctx) and approver_kind(ctx, o) == "head"):
        o["extra"]["approver"] = ctx.author()                 # кто сам согласует — согласован сразу
        _hist(ctx, o, f"согласован сразу ({money(t)})")
        _save(ctx, o, status="ok")
        return [Reply(text=f"✅ Заказ №{o['id']} согласован — можно отправлять поставщику."), order(ctx, oid)]
    o["extra"]["note"] = ""
    _hist(ctx, o, f"на согласование ({money(t)})")
    _save(ctx, o, status="wait")
    who = "директору" if approver_kind(ctx, o) == "owner" else "начальнику закупок"
    out = [Reply(text=f"📨 Заказ №{o['id']} на {money(t)} ушёл на согласование {who}.")]
    out += _notify(ctx, _approvers(ctx, o), f"🛒 Заказ №{o['id']} · {o['supplier']} на {money(t)} ждёт согласования — "
                                            f"от {ctx.author()}.", o["id"])
    out.append(order(ctx, oid))
    return out


@screen("soa", module=M)
def order_approve(ctx, oid):
    o = _order(ctx, oid)
    if not o or o["status"] != "wait" or not can_approve(ctx, o):
        return Reply(text="Согласовывать нечего.")
    o["extra"]["approver"] = ctx.author()
    _hist(ctx, o, "согласован")
    _save(ctx, o, status="ok")
    ctx.db.audit(ctx.author(), "stock_order", f"№{o['id']} {o['supplier']} согласован")
    author = ctx.db.employee(o["emp_id"]) if o.get("emp_id") else None
    out = [Reply(text=f"✅ Заказ №{o['id']} согласован.")]
    out += _notify(ctx, [author] if author else [], f"✅ Заказ №{o['id']} · {o['supplier']} согласован — "
                                                   "отправляйте поставщику.", o["id"])
    out.append(approvals(ctx))
    return out


@screen("sob", module=M)
def order_back(ctx, oid):
    o = _order(ctx, oid)
    if not o or o["status"] != "wait" or not can_approve(ctx, o):
        return Reply(text="Возвращать нечего.")
    return ctx.ask("stock_back", "Что поправить? Комментарий уйдёт менеджеру.", {"id": o["id"]})


@text_input("stock_back")
def on_back(ctx, text, data):
    o = _order(ctx, data["id"])
    note = (text or "").strip()[:300]
    o["extra"]["note"] = note
    _hist(ctx, o, f"возвращён: {note}")
    _save(ctx, o, status="back")
    author = ctx.db.employee(o["emp_id"]) if o.get("emp_id") else None
    out = _notify(ctx, [author] if author else [], f"↩️ Заказ №{o['id']} · {o['supplier']} вернули: {note}", o["id"])
    return out + [approvals(ctx)]


@screen("sos", module=M)
def order_sent(ctx, oid):
    o = _order(ctx, oid)
    if not ok(ctx) or not o or o["status"] != "ok":
        return Reply(text="Отправить поставщику можно согласованный заказ.")
    _hist(ctx, o, "отправлен поставщику")
    _save(ctx, o, status="sent", sent_ts=ctx.today.isoformat())
    ctx.db.audit(ctx.author(), "stock_order", f"№{o['id']} {o['supplier']} отправлен")
    out = [Reply(text=f"🚚 Заказ №{o['id']} в пути — бот вычтет его из «нужно» до приёмки на складе (или {TRANSIT_DAYS} дней).")]
    wh = [e for e in ctx.db.employees() if ctx.cfg.dept(e["dept"]) and ctx.cfg.dept(e["dept"]).kind == "warehouse"
          and e["role"] == "head"]
    out += _notify(ctx, wh, f"🚚 Едет поставка: заказ №{o['id']} · {o['supplier']}, позиций {len(o['items'])}.", o["id"])
    out.append(order(ctx, oid))
    return out


@screen("sow", module=M)
def approvals(ctx):
    if not boss(ctx):
        return Reply(text="Согласует начальник закупок или директор.")
    s = load(ctx)
    ids = waiting(ctx)
    v = View("✅ Заказы на согласовании", f"ждут вас: {len(ids)} · порог директора {money(limit(ctx))}")
    cards = []
    for oid in ids:
        o = _order(ctx, oid)
        t = total(s, o)
        cards.append(Card(f"№{o['id']} · {o['supplier']}", money(t), [f"позиций {len(o['items'])} · оформил {o['author']}"],
                          color="warn" if t > limit(ctx) else None))
    v.cards(cards)
    if not ids:
        v.note("Всё согласовано 👍")
    return Reply(view=v, buttons=rows([Button(f"№{oid}", f"sko:{oid}") for oid in ids[:12]], 4))


@screen("sol", module=M)
def orders(ctx, show="all"):
    if not ok(ctx):
        return Reply(text="Нет доступа.")
    rs = [o for o in (_order(ctx, r["id"]) for r in ctx.db.q("SELECT id FROM stock_orders ORDER BY id DESC LIMIT 60"))
          if mine(ctx, o["supplier"])]
    if show != "all":
        rs = [o for o in rs if o["status"] == show]
    s = load(ctx)
    v = View("🛒 Заказы поставщикам", f"{'все' if show == 'all' else ST[show]} · {len(rs)}")
    v.table(["№", "Поставщик", "Статус", "≈ ₽"], [Row([str(o["id"]), o["supplier"][:24], ST[o["status"]], money(total(s, o), "")],
                                                 sub=o["ts"][:10]) for o in rs[:20]]) \
        if rs else v.note("Заказов нет. Откройте поставщика в «📦 Склад» → «🛒 Заказ».")
    kb = rows([Button(f"№{o['id']}", f"sko:{o['id']}") for o in rs[:12]], 4)
    kb.append([Button("⏳ На согласовании", "sol:wait"), Button("🚚 В пути", "sol:sent")] if show == "all" else
              [Button("Все заказы", "sol")])
    kb.append([Button("‹ Склад", "sk")])
    return Reply(view=v, buttons=kb)


@screen("sox", module=M)
def order_xlsx(ctx, oid):
    o = _order(ctx, oid)
    if not ok(ctx) or not o:
        return Reply(text="Нет доступа.")
    from openpyxl import Workbook
    from openpyxl.styles import Font
    s = load(ctx)
    wb = Workbook()
    ws = wb.active
    ws.title = "Заказ"
    ws.append([f"Заказ №{o['id']} · {o['supplier']} · {ctx.cfg.name} · {ctx.today:%d.%m.%Y}"])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append([])
    ws.append(["Код", "Товар", "Количество", "Ед."])
    for c in ws[3]:
        c.font = Font(bold=True)
    for code, q in o["items"]:
        it = s.items.get(code, {"name": code, "unit": "шт"})
        ws.append([code, it["name"], q, _unit(it)])
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 12, 50
    ws2 = wb.create_sheet("Для нас")
    ws2.append(["Код", "Товар", "Остаток", "Продаётся в месяц", "Дней запаса", "В заказе"])
    for code, q in o["items"]:
        if code in s.items:
            d = s.days_left(code)
            ws2.append([code, s.items[code]["name"], s.items[code]["qty"], round(s.speed(code) * 30),
                        round(d) if d is not None else None, q])
    buf = io.BytesIO()
    wb.save(buf)
    return Reply(text=f"📥 Заказ №{o['id']}", files=[(f"Заказ {o['id']} {o['supplier'][:30]}.xlsx", buf.getvalue())])


# ---------- поиск ----------

@free_text(module=M, order=30)
def find(ctx, text):
    t = (text or "").strip()
    if not ok(ctx) or len(t) < 3 or len(t) > 40:
        return None
    s = load(ctx)
    if not s.items:
        return None
    if t in s.items:
        v = View("🔎 Товар", s.items[t]["supplier"])
        v.cards([_item_card(s, t)])
        return Reply(view=v, buttons=[[Button("‹ Склад", "sk")]])
    if t.isdigit():
        return None
    found = [c for c, it in s.items.items() if t.lower() in it["name"].lower()]
    if not found:
        return None
    v = View(f"🔎 «{t}»", f"найдено {len(found)}")
    v.cards([_item_card(s, c) for c in found[:10]])
    return Reply(view=v, buttons=[[Button("‹ Склад", "sk")]])


# ---------- рассылки ----------

@job("stock_morning", "morning", {"owner", "head", "staff"}, "Утром: что закончилось на складе", module=M)
def morning(ctx):
    if not ok(ctx) or ctx.role == "staff":
        return None
    s = load(ctx)
    if not s.items:
        return None
    out_ = [c for c in s.items if s.cat(c) == "out"]
    low = [c for c in s.items if s.cat(c) == "low"]
    if not out_ and not low:
        return None
    lines = [f"🔴 Закончилось: {len(out_)}", f"🟠 Скоро закончится: {len(low)}"]
    lines += [f"• {s.items[c]['name']}" for c in sorted(out_, key=lambda c: -s.speed(c))[:5]]
    return Reply(text="📦 Склад сегодня\n" + "\n".join(lines), buttons=[[Button("📦 Склад", "sk")]])


# ---------- единицы для зарплаты закупщика: «закончилось» и «залежалось» у своих поставщиков ----------

def _cats(pay):
    cache = pay.__dict__.get("_stock_cats")
    if cache is None:
        from types import SimpleNamespace
        today = min(pay.cal.today, pay.cal.end)
        st = Stock(SimpleNamespace(db=pay.db, cfg=pay.cfg, today=today))
        cache = {}
        for code, it in st.items.items():
            c = st.cat(code)
            if c:
                cache.setdefault(c, []).append(it["supplier"] or "Без поставщика")
        pay._stock_cats = cache
    return cache


def _count(cat):
    def fn(pay, emp):
        own = owners(pay.db)
        sups = _cats(pay).get(cat, [])
        if emp.get("role") == "head" or not own:
            return len(sups)
        return sum(1 for s in sups if own.get(s) == emp["id"])
    return fn


from core.pay import components as _C  # noqa: E402

_C.unit("stockout", _count("out"))
_C.unit("stale", _count("stale"))
