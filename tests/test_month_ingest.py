"""Цикл месяца, правки, ведомость; приём файлов через переводчик клиента."""
import io
import unittest
from datetime import date

from clients.demo.seed import seed
from core.config import load_client
from core.data.db import DB
from core.ingest import excel
from core.ingest.pipeline import ingest
from core.kpi.engine import Kpi
from core.month import cycle
from core.pay.engine import Payroll

TODAY = date(2026, 10, 27)


class Month(unittest.TestCase):
    def setUp(self):
        self.cfg = load_client("demo")
        self.db = DB.open(":memory:")
        seed(self.db, self.cfg, "auto", today=TODAY)
        self.emp = next(e for e in self.db.employees() if e["role"] == "rep")

    def pay(self, month="2026-10"):
        return Payroll(self.cfg, self.db, Kpi(self.cfg, self.db, month, TODAY))

    def test_edit_needs_reason_and_changes_pay(self):
        with self.assertRaises(cycle.MonthError):
            cycle.propose(self.db, "2026-10", self.emp["id"], 5000, "", "РОП")
        before = self.pay().calc(self.emp).total
        eid = cycle.propose(self.db, "2026-10", self.emp["id"], 5000, "Подменял коллегу", "РОП")
        self.assertEqual(self.pay().calc(self.emp).total, before)          # ждёт решения — не считается
        cycle.decide(self.db, eid, True, "Директор")
        self.assertEqual(self.pay().calc(self.emp).total, before + 5000)
        with self.assertRaises(cycle.MonthError):
            cycle.decide(self.db, eid, False, "Директор")                    # повторно решить нельзя

    def test_approve_locks_month_and_keeps_snapshot(self):
        cycle.advance(self.db, "2026-10", "РОП")
        self.assertEqual(cycle.status(self.db, "2026-10"), "review")
        snap = cycle.approve(self.db, "2026-10", "Директор", self.pay())
        self.assertEqual(cycle.status(self.db, "2026-10"), "approved")
        self.assertEqual(cycle.snapshot(self.db, "2026-10"), snap)
        self.assertFalse(cycle.edits(self.db, "2026-10", "wait"))          # нерешённые правки отклонены
        with self.assertRaises(cycle.MonthError):
            cycle.propose(self.db, "2026-10", self.emp["id"], 1, "поздно", "РОП")
        with self.assertRaises(cycle.MonthError):
            cycle.reopen(self.db, "2026-10", "Директор", "")
        cycle.reopen(self.db, "2026-10", "Директор", "ошибка в табеле")
        self.assertEqual(cycle.status(self.db, "2026-10"), "open")

    def test_registry_xlsx(self):
        from openpyxl import load_workbook
        p = self.pay()
        slips = [p.calc(e) for e in p.kpi.employees.values()]
        wb = load_workbook(io.BytesIO(cycle.registry_xlsx(self.cfg, "2026-10", slips)))
        ws = wb["Ведомость"]
        total = ws.cell(ws.max_row, 7).value
        self.assertEqual(total, sum(s.total for s in slips if s.rows))
        self.assertIn("Расчёт по строкам", wb.sheetnames)


def sales_report(month_line, rows):
    """Отчёт «Продажи» в виде, который понимает образец переводчика из clients/_template."""
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    for r in [["Продажи"], [month_line], ["Показатели: Сумма"], ["Группировки строк: Агент, Контрагент, Группа"], []] + rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


class Ingest(unittest.TestCase):
    def setUp(self):
        self.db = DB.open(":memory:")
        self.db.add_employee("Иванов Пётр", "sales", role="rep", keys=["иванов петр"])
        self.file = sales_report("Период: 01.10.2026 - 26.10.2026", [
            ["Б2 ТП Иванов Пётр Сергеевич", 1500], ["  ИП Ромашка (код 1043)", 1500], ["      Сухие смеси", 1000],
            ["      Кровля", 500], ["Сидоров Олег", 300], ["  ООО Лето (код 77)", 300], ["      Кровля", 300]])

    def test_helpers(self):
        self.assertEqual(excel.person_key("Б2 ТП Иванов Пётр Сергеевич"), "иванов петр")
        self.assertEqual(excel.detect_period("Период: Октябрь 2026 г."), ("2026-10", date(2026, 10, 31)))
        self.assertEqual(excel.detect_period("Отбор: Дата > 30.09.2026 23:59:59 И < 01.11.2026 0:00:00")[0], "2026-10")
        self.assertEqual(excel.num("1 234,50"), 1234.5)

    def test_template_parser_and_rules(self):
        res = ingest("_template", self.db, [("продажи.xlsx", self.file)])
        self.assertTrue(res[0].ok, res[0].text())
        self.assertIn("Сидоров Олег", res[0].text())              # человека нет в списке — сказали
        self.assertEqual(sum(r["amount"] for r in self.db.sales("2026-10-01", "2026-10-31")), 1800)
        again = ingest("_template", self.db, [("продажи (копия).xlsx", self.file)])
        self.assertIn("уже принят", again[0].text())                 # тот же файл — не задваиваем
        self.assertEqual(ingest("_template", self.db, [("x.xlsx", sales_report("Что-то", []))])[0].ok, [])

    def test_month_status_blocks_auto_files(self):
        self.db.set_month_status("2026-10", "review", "test")
        self.assertIn("на согласовании", ingest("_template", self.db, [("a.xlsx", self.file)])[0].text())
        self.assertTrue(ingest("_template", self.db, [("a.xlsx", self.file)], manual=True)[0].ok)
        self.db.set_month_status("2026-10", "approved", "test")
        f2 = self.file + b""                                          # другой отпечаток не нужен: месяц закрыт
        self.assertIn("утверждён", ingest("_template", self.db, [("b.xlsx", f2)], manual=True)[0].text())


if __name__ == "__main__":
    unittest.main()
